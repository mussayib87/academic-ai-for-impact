import csv
import hashlib
import io
import json
import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.app.models import (
    AcademicDataImportBatch,
    AcademicDataImportRow,
    AcademicMomentum,
    AcademicRiskOutcome,
    Assessment,
    Attendance,
    AuditLog,
    Course,
    LearningActivity,
    Mark,
    StudentSubjectAcademicRecord,
    StudentSubjectAcademicRecordHistory,
    Student,
    Subject,
)
from backend.app.services.academic_records import normalize_academic_period

MAX_IMPORT_ROWS = 20_000
SCHEMAS = {
    'marks': {'required': {'student_id', 'assessment', 'score', 'recorded_at'}, 'optional': {'maximum_marks', 'course_code'}},
    'assessments': {'required': {'student_id', 'assessment', 'score', 'recorded_at', 'maximum_marks'}, 'optional': {'course_code'}},
    'attendance': {'required': {'student_id', 'attendance_rate', 'observed_at'}, 'optional': {'attended', 'course_code'}},
    'learning_activities': {'required': {'student_id', 'activity_type', 'quantity', 'occurred_at'}, 'optional': {'course_code'}},
    'subject_records': {'required': {'student_id', 'subject_code', 'academic_period'}, 'optional': {'fa1', 'fa1_observed_at', 'fa2', 'fa2_observed_at', 'assignment_mark', 'assignment_observed_at', 'attendance_rate', 'attendance_observed_at'}},
    'risk_outcomes': {'required': {'student_id', 'academic_period', 'risk_target', 'observed_at', 'outcome_source'}, 'optional': set()},
}
HEADER_ALIASES = {
    'usn': 'student_id',
    'student_identifier': 'student_id',
    'subject': 'subject_code',
    'academicperiod': 'academic_period',
    'semester': 'academic_period',
    'term': 'academic_period',
    'fa-1': 'fa1',
    'fa-2': 'fa2',
    'assignment': 'assignment_mark',
}


class ImportValidationError(ValueError):
    def __init__(self, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.errors = errors or []


def _normalized_header(value: Any) -> str:
    header = str(value or '').strip().lower().replace(' ', '_')
    return HEADER_ALIASES.get(header, header)


def _read_rows(filename: str, content: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == '.csv':
            text = content.decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(text, newline=''))
            if not reader.fieldnames:
                raise ImportValidationError('The CSV file is empty or has no header row.')
            headers = [_normalized_header(value) for value in reader.fieldnames]
            if any(not header for header in headers):
                raise ImportValidationError('The file contains an empty column name.')
            if len(headers) != len(set(headers)):
                raise ImportValidationError('The file contains duplicate column names.')
            rows = []
            for row_number, row in enumerate(reader, start=2):
                if None in row:
                    raise ImportValidationError(f'CSV row {row_number} has more values than the header row.')
                if len(rows) >= MAX_IMPORT_ROWS:
                    raise ImportValidationError(f'Files may contain at most {MAX_IMPORT_ROWS} data rows.')
                rows.append({headers[index]: row.get(reader.fieldnames[index]) for index in range(len(headers))})
        elif suffix == '.xlsx':
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = workbook.active
                values = sheet.iter_rows(values_only=True)
                headers_row = next(values, None)
                if not headers_row:
                    raise ImportValidationError('The Excel sheet is empty or has no header row.')
                headers = [_normalized_header(value) for value in headers_row]
                if any(not header for header in headers):
                    raise ImportValidationError('The file contains an empty column name.')
                if len(headers) != len(set(headers)):
                    raise ImportValidationError('The file contains duplicate column names.')
                rows = []
                for row_number, row in enumerate(values, start=2):
                    if not any(value is not None for value in row):
                        continue
                    if len(rows) >= MAX_IMPORT_ROWS:
                        raise ImportValidationError(f'Files may contain at most {MAX_IMPORT_ROWS} data rows.')
                    rows.append(dict(zip(headers, row)))
            finally:
                workbook.close()
        else:
            raise ImportValidationError('Upload a .csv or .xlsx academic data file.')
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError(f'Could not read the uploaded file: {exc}') from exc
    if len(rows) > MAX_IMPORT_ROWS:
        raise ImportValidationError(f'Files may contain at most {MAX_IMPORT_ROWS} data rows.')
    if not rows:
        raise ImportValidationError('The file contains no data rows.')
    return rows


def _normalized_row(row: dict[str, Any]) -> dict[str, str]:
    return {_normalized_header(key): str(value).strip() if value is not None else '' for key, value in row.items()}


def _row_checksum(dataset_type: str, row: dict[str, Any]) -> str:
    normalized = _normalized_row(row)
    for identifier in ('student_id', 'subject_code', 'academic_period', 'risk_target'):
        if identifier in normalized:
            normalized[identifier] = normalized[identifier].casefold()
    data = json.dumps([dataset_type, normalized], sort_keys=True, separators=(',', ':'), ensure_ascii=False)
    return hashlib.sha256(data.encode('utf-8')).hexdigest()


def _number(value: Any, column: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'{column} must be numeric.') from exc
    if not math.isfinite(result):
        raise ValueError(f'{column} must be a finite number.')
    return result


def _date(value: Any, column: str) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    try:
        parsed = datetime.fromisoformat(str(value).strip().replace('Z', '+00:00'))
        return parsed.replace(tzinfo=None)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'{column} must be an ISO date or datetime.') from exc


def _audit_values(values: dict[str, Any]) -> dict[str, Any]:
    return {key: value.isoformat() if isinstance(value, datetime) else value for key, value in values.items()}


def _boolean(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {'true', '1', 'yes', 'present', 'attended'}:
        return True
    if normalized in {'false', '0', 'no', 'absent', 'missed'}:
        return False
    raise ValueError('attended must be true/false, yes/no, present/absent, or 1/0.')


def _students_by_identifier(db: Session) -> dict[str, list[Student]]:
    result: dict[str, list[Student]] = {}
    for student in db.query(Student).all():
        result.setdefault(student.student_id.strip().casefold(), []).append(student)
    return result


def import_academic_data(db: Session, *, user_id: int, filename: str, content: bytes, dataset_type: str) -> dict[str, Any]:
    kind = dataset_type.strip().lower()
    if kind not in SCHEMAS:
        raise ImportValidationError(f'Unsupported dataset type. Choose one of: {", ".join(SCHEMAS)}.')
    rows = _read_rows(filename, content)
    headers = set(_normalized_row(rows[0]))
    missing = sorted(SCHEMAS[kind]['required'] - headers)
    if missing:
        raise ImportValidationError(
            f'Missing required columns for {kind}: {", ".join(missing)}.',
            [{'row': 1, 'errors': [f'Missing required column: {column}.' for column in missing]}],
        )

    checksum = hashlib.sha256(content).hexdigest()
    previous_batch = db.query(AcademicDataImportBatch).filter_by(dataset_type=kind, file_checksum=checksum).first()
    if previous_batch:
        return {
            'batch_id': previous_batch.id,
            'dataset_type': kind,
            'source_filename': previous_batch.source_filename,
            'imported_rows': 0,
            'rejected_rows': 0,
            'unmatched_students': 0,
            'duplicate_rows': len(rows),
            'unmatched_student_ids': [],
            'validation_errors': [],
            'duplicate_import': True,
        }

    students = _students_by_identifier(db)
    subjects = {subject.code.strip().casefold(): subject for subject in db.query(Subject).all()}
    batch = AcademicDataImportBatch(
        dataset_type=kind,
        source_filename=Path(filename).name[:255],
        file_checksum=checksum,
        imported_by_user_id=user_id,
        validation_errors=[],
    )
    db.add(batch)
    db.flush()

    accepted: list[tuple[dict[str, Any], Student, str, int]] = []
    row_records: list[AcademicDataImportRow] = []
    validation_errors: list[dict[str, Any]] = []
    row_records_by_number: dict[int, AcademicDataImportRow] = {}
    unmatched: set[str] = set()
    seen_checksums: set[str] = set()
    duplicates = 0
    for row_number, raw_row in enumerate(rows, start=2):
        row = _normalized_row(raw_row)
        row_hash = _row_checksum(kind, raw_row)
        if row_hash in seen_checksums or db.query(AcademicDataImportRow.id).filter_by(dataset_type=kind, row_checksum=row_hash).first():
            duplicates += 1
            continue
        seen_checksums.add(row_hash)
        errors: list[str] = []
        identifier = row.get('student_id', '').strip()
        matches = students.get(identifier.casefold(), [])
        student = matches[0] if len(matches) == 1 else None
        if not identifier:
            errors.append('student_id is required.')
        elif not matches:
            errors.append(f'Unknown student_id: {identifier}.')
            unmatched.add(identifier)
        elif len(matches) > 1:
            errors.append(f'Ambiguous student_id: {identifier}.')

        values: dict[str, Any] = {}
        try:
            if kind == 'subject_records':
                subject_code = row.get('subject_code', '').strip()
                subject = subjects.get(subject_code.casefold())
                if not subject_code:
                    raise ValueError('subject_code is required.')
                if subject is None:
                    raise ValueError(f'Unknown subject_code: {subject_code}.')
                period = normalize_academic_period(row.get('academic_period', ''))
                updates: dict[str, Any] = {}
                for field, maximum in (('fa1', 25), ('fa2', 25), ('assignment_mark', 25), ('attendance_rate', 100)):
                    raw_value = row.get(field, '').strip()
                    if raw_value:
                        value = _number(raw_value, field)
                        if not 0 <= value <= maximum:
                            raise ValueError(f'{field} must be between 0 and {maximum}.')
                        updates[field] = value
                for field in ('fa1_observed_at', 'fa2_observed_at', 'assignment_observed_at', 'attendance_observed_at'):
                    raw_date = row.get(field, '').strip()
                    if raw_date:
                        updates[field] = _date(raw_row.get(field), field)
                if not updates:
                    raise ValueError('At least one of fa1, fa2, assignment_mark, or attendance_rate is required.')
                values = {'subject': subject, 'academic_period': period, 'updates': updates}
            elif kind == 'risk_outcomes':
                period = normalize_academic_period(row.get('academic_period', ''))
                risk_target = row.get('risk_target', '').strip().upper()
                if risk_target not in {'HIGH', 'MEDIUM', 'LOW'}:
                    raise ValueError('risk_target must be HIGH, MEDIUM, or LOW.')
                outcome_source = row.get('outcome_source', '').strip()
                if not outcome_source:
                    raise ValueError('outcome_source is required and must identify an independent, non-model outcome source.')
                normalized_outcome_source = re.sub(r'[^a-z0-9]+', ' ', outcome_source.casefold()).strip()
                if any(marker in normalized_outcome_source for marker in (
                    'model prediction', 'model generated', 'shap', 'inference', 'synthetic',
                    'ai generated label', 'risk model', 'classifier output',
                    'derived from marks', 'automated prediction',
                )):
                    raise ValueError('risk outcomes cannot be sourced from model predictions, SHAP, or labels derived from these marks.')
                values = {'academic_period': period, 'risk_target': risk_target, 'observed_at': _date(raw_row.get('observed_at'), 'observed_at'), 'outcome_source': outcome_source}
            elif kind in {'marks', 'assessments'}:
                title = str(row.get('assessment', '')).strip()
                if not title:
                    raise ValueError('assessment is required.')
                score = _number(row.get('score'), 'score')
                maximum = _number(row.get('maximum_marks', 100), 'maximum_marks')
                if maximum <= 0 or score < 0 or score > maximum:
                    raise ValueError('score must be between 0 and maximum_marks, and maximum_marks must be positive.')
                recorded_at = _date(raw_row.get('recorded_at'), 'recorded_at')
                course_code = str(row.get('course_code', '')).strip()
                course = db.query(Course).filter(Course.code == course_code).first() if course_code else None
                if course_code and course is None:
                    raise ValueError(f'Unknown course_code: {course_code}.')
                if not course and student:
                    course = db.query(Course).filter(Course.id == student.course_id).first() if student.course_id else None
                if student and student.course_id and course and student.course_id != course.id:
                    raise ValueError("course_code does not match the student's linked course.")
                if student and not course:
                    raise ValueError('Student has no course; provide an existing course_code.')
                assessment = db.query(Assessment).filter(Assessment.course_id == course.id, Assessment.title == title).first() if course else None
                if assessment and assessment.maximum_marks != maximum:
                    raise ValueError('maximum_marks does not match the existing assessment; existing assessments are not overwritten.')
                values = {'score': score, 'maximum': maximum, 'recorded_at': recorded_at, 'course': course, 'assessment': assessment, 'title': title}
            elif kind == 'attendance':
                rate = _number(row.get('attendance_rate'), 'attendance_rate')
                if not 0 <= rate <= 100:
                    raise ValueError('attendance_rate must be between 0 and 100.')
                observed_at = _date(raw_row.get('observed_at'), 'observed_at')
                attended = _boolean(raw_row.get('attended')) if row.get('attended') else rate > 0
                course_code = str(row.get('course_code', '')).strip()
                course = db.query(Course).filter(Course.code == course_code).first() if course_code else None
                if course_code and course is None:
                    raise ValueError(f'Unknown course_code: {course_code}.')
                if not course and student:
                    course = db.query(Course).filter(Course.id == student.course_id).first() if student.course_id else None
                values = {'rate': rate, 'observed_at': observed_at, 'attended': attended, 'course': course}
            else:
                activity_type = str(row.get('activity_type', '')).strip()
                if not activity_type:
                    raise ValueError('activity_type is required.')
                quantity = _number(row.get('quantity'), 'quantity')
                if quantity < 0:
                    raise ValueError('quantity must not be negative.')
                occurred_at = _date(raw_row.get('occurred_at'), 'occurred_at')
                course_code = str(row.get('course_code', '')).strip()
                course = db.query(Course).filter(Course.code == course_code).first() if course_code else None
                if course_code and course is None:
                    raise ValueError(f'Unknown course_code: {course_code}.')
                if not course and student:
                    course = db.query(Course).filter(Course.id == student.course_id).first() if student.course_id else None
                values = {'activity_type': activity_type, 'quantity': quantity, 'occurred_at': occurred_at, 'course': course}
        except ValueError as exc:
            errors.append(str(exc))

        if errors:
            validation_errors.append({
                'row': row_number,
                'student_id': identifier or None,
                'subject_code': row.get('subject_code') or None,
                'academic_period': row.get('academic_period') or None,
                'errors': errors,
            })
        else:
            accepted.append((values, student, row_hash, row_number))
        import_row = AcademicDataImportRow(
            batch_id=batch.id,
            dataset_type=kind,
            row_checksum=row_hash,
            row_number=row_number,
            status='rejected' if errors else 'imported',
            student_id=student.id if student else None,
            errors=errors,
        )
        row_records.append(import_row)
        row_records_by_number[row_number] = import_row

    conflict_groups: dict[tuple[Any, ...], list[tuple[dict[str, Any], Student, str, int]]] = {}
    if kind == 'subject_records':
        for accepted_row in accepted:
            values, student, _, _ = accepted_row
            key = (student.id, values['subject'].id, values['academic_period'])
            conflict_groups.setdefault(key, []).append(accepted_row)
        conflicting_hashes = set()
        for key, grouped_rows in conflict_groups.items():
            conflicting_fields = [
                field for field in ('fa1', 'fa1_observed_at', 'fa2', 'fa2_observed_at', 'assignment_mark', 'assignment_observed_at', 'attendance_rate', 'attendance_observed_at')
                if len({row[0]['updates'][field] for row in grouped_rows if field in row[0]['updates']}) > 1
            ]
            if not conflicting_fields:
                continue
            subject = grouped_rows[0][0]['subject']
            period = grouped_rows[0][0]['academic_period']
            for _, student, row_hash, row_number in grouped_rows:
                message = f'Conflicting duplicate {", ".join(conflicting_fields)} values for {subject.code} in {period}; upload corrections in a later import.'
                import_row = row_records_by_number[row_number]
                import_row.status = 'rejected'
                import_row.errors = [message]
                validation_errors.append({
                    'row': row_number,
                    'student_id': student.student_id,
                    'subject_code': subject.code,
                    'academic_period': period,
                    'errors': [message],
                })
                conflicting_hashes.add(row_hash)
        if conflicting_hashes:
            accepted = [row for row in accepted if row[2] not in conflicting_hashes]
    elif kind == 'risk_outcomes':
        outcome_groups: dict[tuple[Any, ...], list[tuple[dict[str, Any], Student, str, int]]] = {}
        for accepted_row in accepted:
            values, student, _, _ = accepted_row
            outcome_groups.setdefault((student.id, values['academic_period']), []).append(accepted_row)
        conflicting_hashes = set()
        for (_, period), grouped_rows in outcome_groups.items():
            targets = {row[0]['risk_target'] for row in grouped_rows}
            if len(targets) < 2:
                continue
            for _, student, row_hash, row_number in grouped_rows:
                message = f'Conflicting risk_target values for {period}; upload corrections in a later import.'
                import_row = row_records_by_number[row_number]
                import_row.status = 'rejected'
                import_row.errors = [message]
                validation_errors.append({'row': row_number, 'student_id': student.student_id, 'academic_period': period, 'errors': [message]})
                conflicting_hashes.add(row_hash)
        if conflicting_hashes:
            accepted = [row for row in accepted if row[2] not in conflicting_hashes]

    created_assessments: dict[tuple[int, str], Assessment] = {}
    for values, student, row_hash, row_number in accepted:
        if kind == 'subject_records':
            subject = values['subject']
            period = values['academic_period']
            record = db.query(StudentSubjectAcademicRecord).filter_by(student_id=student.id, subject_id=subject.id, academic_period=period).first()
            if record is None:
                record = StudentSubjectAcademicRecord(student_id=student.id, subject_id=subject.id, academic_period=period)
                db.add(record)
                db.flush()
            old_values = {field: getattr(record, field) for field in values['updates']}
            for field, value in values['updates'].items():
                setattr(record, field, value)
            record.import_batch_id = batch.id
            record.updated_by_user_id = user_id
            db.flush()
            db.add(StudentSubjectAcademicRecordHistory(
                academic_record_id=record.id,
                import_batch_id=batch.id,
                changed_by_user_id=user_id,
                change_kind='import',
                fa1=record.fa1,
                fa1_observed_at=record.fa1_observed_at,
                fa2=record.fa2,
                fa2_observed_at=record.fa2_observed_at,
                assignment_mark=record.assignment_mark,
                assignment_observed_at=record.assignment_observed_at,
                attendance_rate=record.attendance_rate,
                attendance_observed_at=record.attendance_observed_at,
            ))
            row_records_by_number[row_number].academic_record_id = record.id
            db.add(AuditLog(
                actor_user_id=user_id,
                action='academic_subject_record.imported',
                resource_type='student_subject_academic_record',
                resource_id=f'{student.id}:{subject.code}:{period}',
                detail={'student_id': student.student_id, 'subject_code': subject.code, 'academic_period': period, 'old_values': _audit_values(old_values), 'new_values': _audit_values(values['updates']), 'batch_id': batch.id},
            ))
        elif kind == 'risk_outcomes':
            outcome = db.query(AcademicRiskOutcome).filter_by(student_id=student.id, academic_period=values['academic_period']).first()
            if outcome is None:
                outcome = AcademicRiskOutcome(student_id=student.id, academic_period=values['academic_period'])
                db.add(outcome)
            outcome.risk_level = values['risk_target']
            outcome.observed_at = values['observed_at']
            outcome.outcome_source = values['outcome_source']
            outcome.import_batch_id = batch.id
            outcome.recorded_by_user_id = user_id
            db.add(AuditLog(
                actor_user_id=user_id,
                action='academic_risk_outcome.imported',
                resource_type='academic_risk_outcome',
                resource_id=f'{student.id}:{values["academic_period"]}',
                detail={'student_id': student.student_id, 'academic_period': values['academic_period'], 'risk_target': values['risk_target'], 'outcome_source': values['outcome_source'], 'batch_id': batch.id},
            ))
        elif kind in {'marks', 'assessments'}:
            course = values['course']
            assessment_key = (course.id, values['title'])
            assessment = values['assessment'] or created_assessments.get(assessment_key)
            if assessment is None:
                assessment = Assessment(course_id=course.id, title=values['title'], maximum_marks=values['maximum'], import_batch_id=batch.id)
                db.add(assessment)
                db.flush()
                created_assessments[assessment_key] = assessment
            db.add(Mark(student_id=student.id, assessment_id=assessment.id, score=values['score'], recorded_at=values['recorded_at'], import_batch_id=batch.id))
            db.add(AcademicMomentum(student_id=student.id, metric='assessment', value=values['score'] / values['maximum'] * 100, observed_at=values['recorded_at'], import_batch_id=batch.id))
        elif kind == 'attendance':
            db.add(Attendance(student_id=student.id, course_id=values['course'].id if values['course'] else None, observed_at=values['observed_at'], attended=values['attended'], attendance_rate=values['rate'], import_batch_id=batch.id))
            db.add(AcademicMomentum(student_id=student.id, metric='attendance', value=values['rate'], observed_at=values['observed_at'], import_batch_id=batch.id))
        else:
            db.add(LearningActivity(student_id=student.id, course_id=values['course'].id if values['course'] else None, activity_type=values['activity_type'], quantity=values['quantity'], occurred_at=values['occurred_at'], import_batch_id=batch.id))
            db.add(AcademicMomentum(student_id=student.id, metric='activity', value=values['quantity'], observed_at=values['occurred_at'], import_batch_id=batch.id))

    batch.imported_rows = len(accepted)
    batch.rejected_rows = len(validation_errors)
    batch.duplicate_rows = duplicates
    batch.validation_errors = validation_errors
    db.add_all(row_records)
    try:
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {
        'batch_id': batch.id,
        'dataset_type': kind,
        'source_filename': batch.source_filename,
        'imported_rows': len(accepted),
        'rejected_rows': len(validation_errors),
        'unmatched_students': len(unmatched),
        'unmatched_student_ids': sorted(unmatched),
        'duplicate_rows': duplicates,
        'validation_errors': validation_errors,
        'duplicate_import': False,
    }